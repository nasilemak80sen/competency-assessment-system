import pandas as pd
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from analytics import gap_analysis_individual, gap_summary
from models import Assessment, Base, Personnel
import db_ops


def test_assessment_date_is_nullable_and_add_assessment_accepts_missing_date():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    person = Personnel(staff_id="TEST-UNDATED", name="Undated Person", is_deleted=False)
    session.add(person)
    session.commit()

    ok, _, assessment_id = db_ops.add_assessment(
        session, person.id, {"assessment_date": None, "remarks": "Source date: TBA"}
    )

    assert ok is True
    assessment = session.get(Assessment, assessment_id)
    assert assessment is not None
    assert assessment.assessment_date is None
    assert assessment.remarks == "Source date: TBA"


def test_workbook_gap_convention_counts_positive_values_as_shortfalls():
    frame = pd.DataFrame([
        {"Name": "Shortfall", "G--B1": 1.0, "G--B2": 2.0},
        {"Name": "Meets", "G--B1": 0.0, "G--B2": -1.0},
    ])

    result = gap_summary(frame).set_index("Name")
    assert result.loc["Shortfall", "Gaps Count"] == 2
    assert result.loc["Shortfall", "Gap Status"] == ">1 Gap"
    assert result.loc["Meets", "Gaps Count"] == 0
    assert result.loc["Meets", "Gap Status"] == "No Gap"


def test_future_grade_gap_uses_requirement_minus_actual():
    row = pd.Series({"B1": 2.0})
    result = gap_analysis_individual(
        row,
        ruler_requirements={"P4": {"B1": 3.0}},
        current_sg="P3",
        include_future_grades=True,
    )
    future = result[result["Target Grade"] == "P4"].iloc[0]
    assert future["Gap"] == 1.0
    assert future["Status"] == "Minor Gap"


def test_load_master_data_does_not_invent_full_date_from_year_or_markers(tmp_path):
    import openpyxl
    from data_loader import load_master_data

    path = tmp_path / "assessment_dates.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All"
    headers = ["Staff ID", "SG", "Last Assesment Date", "Staff Position", "Name"]
    for col, value in enumerate(headers, 1):
        ws.cell(3, col, value)
    for row, (sid, name, value) in enumerate([
        ("1", "Year Only", "2021"),
        ("2", "Ongoing", "Ongoing"),
        ("3", "TBA", "TBA"),
        ("4", "Full Date", "2021-05-12"),
    ], 4):
        ws.cell(row, 1, sid)
        ws.cell(row, 5, name)
        ws.cell(row, 3, value)
    wb.save(path)

    df = load_master_data(str(path))
    indexed = df.set_index("Name")
    assert pd.isna(indexed.loc["Year Only", "Last Assesment Date"])
    assert indexed.loc["Year Only", "Last Assesment Date Source"] == "2021"
    assert pd.isna(indexed.loc["Ongoing", "Last Assesment Date"])
    assert indexed.loc["Ongoing", "Last Assesment Date Source"] == "Ongoing"
    assert pd.isna(indexed.loc["TBA", "Last Assesment Date"])
    assert indexed.loc["Full Date", "Last Assesment Date"].date().isoformat() == "2021-05-12"

 
def test_importer_preserves_date_markers_in_assessment_remarks():
    row = pd.Series({
        "Remarks": "Original remark",
        "Last Assesment Date": pd.NaT,
        "Last Assesment Date Source": "TBA",
    })
    assert db_ops._assessment_source_remarks(row) == "Original remark\nSource Last Assesment Date field: TBA"


def test_runtime_database_uses_validated_rebuilt_file():
    import sqlite3

    from config import BASE_DIR, DATABASE_PATH, DATABASE_URL

    expected_path = (BASE_DIR / "re_competency_rebuilt.db").resolve()
    assert DATABASE_PATH == expected_path
    assert DATABASE_URL == f"sqlite:///{expected_path.as_posix()}"
    assert DATABASE_PATH.is_file(), f"Configured database is missing: {DATABASE_PATH}"

    with sqlite3.connect(DATABASE_PATH) as connection:
        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        tables = {
            row[0]
            for row in connection.execute(
                "SELECT name FROM sqlite_master WHERE type = 'table'"
            )
        }
        assert {
            "personnel",
            "assessments",
            "competency_scores",
            "summary_scores",
            "cv_documents",
        }.issubset(tables)
        assert connection.execute("SELECT COUNT(*) FROM personnel").fetchone()[0] > 0

def test_load_ruler_accepts_trailing_underscore_sheet_name(tmp_path):
    import openpyxl

    from config import SCORE_COLS
    from data_loader import load_ruler_and_tech_mapping

    workbook_path = tmp_path / "ruler_alias.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Ruler_"

    # Match the source workbook: competency headers begin in column E,
    # ruler type is column B, SG is column C, and type labels are group headers.
    for column, header in enumerate(SCORE_COLS, start=5):
        sheet.cell(row=4, column=column, value=header)

    rows = [
        (None, "P1", "Junior Executive"),
        (None, "P2", "Executive"),
        ("RDP", "P3", "Senior Executive"),
        (None, "P4", "Senior RE"),
        ("RMS", "P3", "Senior Executive"),
        (None, "P4", "Senior RE"),
        ("RSS", "P3", "Senior Executive"),
        (None, "P4", "Senior RE"),
    ]
    for row_number, (ruler_type, sg, position) in enumerate(rows, start=5):
        if ruler_type is not None:
            sheet.cell(row=row_number, column=2, value=ruler_type)
        sheet.cell(row=row_number, column=3, value=sg)
        sheet.cell(row=row_number, column=4, value=position)
        for column in range(5, 5 + len(SCORE_COLS)):
            sheet.cell(row=row_number, column=column, value=2)

    workbook.save(workbook_path)
    workbook.close()

    ruler_map, _ = load_ruler_and_tech_mapping(str(workbook_path))

    assert {"BASE", "RDP", "RMS", "RSS"}.issubset(ruler_map)
    assert set(ruler_map["BASE"]) == {"P1", "P2"}
    assert set(ruler_map["RDP"]) == {"P3", "P4"}
    assert set(ruler_map["RMS"]) == {"P3", "P4"}
    assert set(ruler_map["RSS"]) == {"P3", "P4"}
    assert ruler_map["RDP"]["P3"]["B1"] == 2.0
    assert ruler_map["RDP"]["P4"]["B1"] == 2.0

def test_resolve_excel_path_prefers_updated_xlsx_over_legacy_xlsm(tmp_path, monkeypatch):
    import config

    monkeypatch.delenv("COMPETENCY_EXCEL_PATH", raising=False)
    monkeypatch.delenv("RE_EXCEL_PATH", raising=False)
    monkeypatch.setattr(config, "BASE_DIR", tmp_path)
    legacy = tmp_path / "RE Fraternity Jul2026_Master.xlsm"
    current = tmp_path / "RE Fraternity Jul2026_Master.xlsx"
    legacy.write_bytes(b"legacy")
    current.write_bytes(b"current")

    assert config.resolve_excel_path() == str(current.resolve())


def test_load_master_data_normalizes_single_hyphen_gap_headers(tmp_path):
    import openpyxl
    from data_loader import load_master_data

    path = tmp_path / "gap_headers.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "All"
    headers = ["Staff ID", "SG", "Last Assesment Date", "Staff Position", "Name", "B1", "R-B1", "G-B1"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=column, value=header)
    sheet.cell(row=4, column=1, value="TEST-GAP")
    sheet.cell(row=4, column=2, value="P3")
    sheet.cell(row=4, column=4, value="Senior Executive")
    sheet.cell(row=4, column=5, value="Gap Test Person")
    sheet.cell(row=4, column=6, value=2)
    sheet.cell(row=4, column=7, value=3)
    sheet.cell(row=4, column=8, value=0.25)
    workbook.save(path)
    workbook.close()

    result = load_master_data(str(path)).iloc[0]
    assert result["B1"] == 2
    assert result["R-B1"] == 3
    assert result["G--B1"] == 0.25
    assert "G-B1" not in result.index


def test_bulk_import_refreshes_existing_assessment_and_scores_from_workbook():
    from datetime import date

    from config import SCORE_COLS
    from models import CompetencyScore
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()

    person = Personnel(
        staff_id="TEST-REIMPORT",
        name="Reimport Person",
        assessment_level="Staff",
        is_active=True,
        is_deleted=False,
    )
    session.add(person)
    session.commit()

    assessment = Assessment(
        personnel_id=person.id,
        assessment_date=date(2025, 5, 12),
        assessment_level="Staff",
        remarks="old workbook value",
    )
    session.add(assessment)
    session.flush()
    session.add(
        CompetencyScore(
            assessment_id=assessment.id,
            personnel_id=person.id,
            competency_code="B1",
            competency_type="B",
            actual_score=1,
            requirement_score=2,
            gap_score=1,
            achievement_pct=50,
        )
    )
    session.commit()

    row = {
        "Staff ID": "TEST-REIMPORT",
        "Name": "Reimport Person",
        "SG": "P3",
        "Staff Position": "Senior Executive",
        "Assessment Level": "UPE",
        "Last Assesment Date": date(2025, 5, 12),
        "Remarks": "latest workbook value",
    }
    for code in SCORE_COLS:
        row[code] = None
        row[f"R-{code}"] = None
        row[f"G--{code}"] = None
    row["B1"] = 2
    row["R-B1"] = 3
    row["G--B1"] = 0.25  # Source value wins over recomputed 1.0.

    result = db_ops.bulk_import_from_df(session, pd.DataFrame([row]))
    assert result["added"] == 0
    assert result["updated"] == 1
    assert result["errors"] == 0

    refreshed_person = session.query(Personnel).filter_by(staff_id="TEST-REIMPORT").one()
    refreshed_assessment = session.query(Assessment).filter_by(personnel_id=person.id).one()
    refreshed_score = (
        session.query(CompetencyScore)
        .filter_by(assessment_id=refreshed_assessment.id, competency_code="B1")
        .one()
    )
    assert refreshed_person.assessment_level == "UPE"
    assert refreshed_assessment.assessment_level == "UPE"
    assert refreshed_assessment.remarks == "latest workbook value"
    assert refreshed_score.actual_score == 2
    assert refreshed_score.requirement_score == 3
    assert refreshed_score.gap_score == 0.25
    assert session.query(CompetencyScore).filter_by(assessment_id=refreshed_assessment.id).count() == len(SCORE_COLS)

    session.close()
    engine.dispose()

def test_assessment_level_values_match_workbook_and_normalize_case(tmp_path):
    import openpyxl
    from config import ASSESSMENT_LEVELS
    from data_loader import load_master_data

    assert ASSESSMENT_LEVELS == [
        "Staff", "UPE", "Principal", "TPCP", "Custodian", "Research",
        "TBA", "N/A", "Ongoing",
    ]

    path = tmp_path / "assessment_levels.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "All"
    headers = ["Staff ID", "SG", "Last Assesment Date", "Staff Position", "Name", "Assessment Level"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=column, value=header)
    records = [
        ("A1", "Person Staff", "staff"),
        ("A2", "Person UPE", "UPE"),
        ("A3", "Person TBA", "TBA"),
        ("A4", "Person Custodian", "Custodian"),
    ]
    for row_number, (staff_id, name, level) in enumerate(records, start=4):
        sheet.cell(row=row_number, column=1, value=staff_id)
        sheet.cell(row=row_number, column=2, value="P3")
        sheet.cell(row=row_number, column=4, value="Senior Executive")
        sheet.cell(row=row_number, column=5, value=name)
        sheet.cell(row=row_number, column=6, value=level)
    workbook.save(path)
    workbook.close()

    loaded = load_master_data(str(path)).set_index("Name")
    assert loaded.loc["Person Staff", "Assessment Level"] == "Staff"
    assert loaded.loc["Person UPE", "Assessment Level"] == "UPE"
    assert loaded.loc["Person TBA", "Assessment Level"] == "TBA"
    assert loaded.loc["Person Custodian", "Assessment Level"] == "Custodian"

